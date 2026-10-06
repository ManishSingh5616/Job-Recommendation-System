import io

import pdfplumber
from docx import Document


def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join((page.extract_text() or "") for page in pdf.pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
    else:
        raise ValueError("Only PDF and DOCX files are supported.")

    text = text.strip()
    if len(text) < 50:
        raise ValueError(
            "No readable text found. The file may be a scanned image; upload a text-based PDF or DOCX."
        )
    return text
